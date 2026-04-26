import json
from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse, HttpResponse
from django.db.models import Sum, Count, Q
from django.utils import timezone
from datetime import timedelta
from .models import * # Barcha modellarni (Category, Product, Branch, Order, OrderItem) import qilamiz
from .forms import CategoryForm, ProductForm, BranchForm, OrderForm

# Activity logging helper
def log_action(request, action, entity, entity_id, message):
    try:
        AuditLog.objects.create(
            user=request.user if request.user.is_authenticated else None,
            action=action,
            entity=entity,
            entity_id=entity_id,
            message=message,
            ip_address=request.META.get('REMOTE_ADDR'),
            user_agent=request.META.get('HTTP_USER_AGENT')
        )
    except:
        pass # Logging failure shouldn't stop the main operation

# Client Views
def product_list(request):
    categories = Category.objects.all()
    products = Product.objects.all()
    branches = Branch.objects.all()
    
    # Session based cart
    cart = request.session.get('cart', {})
    cart_total = 0
    for item in cart.values():
        if isinstance(item, dict):
            cart_total += item.get('price', 0) * item.get('quantity', 1)
        
    context = {
        'categories': categories,
        'products': products,
        'branches': branches,
        'cart_total': cart_total,
        'cart_count': len(cart)
    }
    return render(request, 'index_1.html', context)

def filiallar_page(request):
    branches = Branch.objects.all()
    return render(request, 'filiallar.html', {'branches': branches})

# Auth
from django.contrib.auth import authenticate, login

def login_page(request):
    if request.method == 'POST':
        u = request.POST.get('username')
        p = request.POST.get('password')
        user = authenticate(username=u, password=p)
        if user is not None:
            login(request, user)
            log_action(request, 'login', 'User', user.id, f"Admin panelga kirish: {user.username}")
            return redirect('dashboard')
        else:
            return render(request, 'login.html', {'error': "Login yoki parol noto'g'ri!"})
    return render(request, 'login.html')

def logout_page(request):
    from django.contrib.auth import logout
    logout(request)
    return redirect('login_page')

# Cart & Order Logic
def add_to_cart(request):
    if request.method == 'POST':
        p_id = request.POST.get('product_id')
        product = get_object_or_404(Product, id=p_id)
        cart = request.session.get('cart', {})
        
        if str(p_id) in cart:
            cart[str(p_id)]['quantity'] += 1
        else:
            cart[str(p_id)] = {
                'name': product.name,
                'price': float(product.price),
                'quantity': 1,
                'image': product.image.url if product.image else ''
            }
        
        request.session['cart'] = cart
        total = sum(item.get('price', 0) * item.get('quantity', 1) for item in cart.values() if isinstance(item, dict))
        return JsonResponse({'success': True, 'cart_total': total})
    return JsonResponse({'success': False})

def cart_modal(request):
    cart = request.session.get('cart', {})
    items = []
    total = 0
    for p_id, item in cart.items():
        subtotal = item['price'] * item['quantity']
        items.append({
            'id': p_id,
            'name': item['name'],
            'price': item['price'],
            'quantity': item['quantity'],
            'image': item.get('image', ''),
            'total': subtotal
        })
        total += subtotal
        
    return render(request, 'cart_modal.html', {'items': items, 'total': total})

def update_cart(request):
    if request.method == 'POST':
        p_id = request.POST.get('product_id')
        action = request.POST.get('action') # 'add' or 'remove'
        cart = request.session.get('cart', {})
        
        if p_id in cart:
            if action == 'add':
                cart[p_id]['quantity'] += 1
            elif action == 'remove':
                if cart[p_id]['quantity'] > 1:
                    cart[p_id]['quantity'] -= 1
                else:
                    del cart[p_id]
            request.session['cart'] = cart
            total = sum(item.get('price', 0) * item.get('quantity', 1) for item in cart.values() if isinstance(item, dict))
            return JsonResponse({'success': True, 'cart_total': total})
    return JsonResponse({'success': False})

def clear_cart(request):
    request.session['cart'] = {}
    return JsonResponse({'success': True})

def order_page(request):
    cart = request.session.get('cart', {})
    if not cart:
        return redirect('product_list')
        
    items = []
    total = 0
    for p_id, item in cart.items():
        subtotal = item['price'] * item['quantity']
        items.append({
            'name': item['name'],
            'price': item['price'],
            'quantity': item['quantity'],
            'total': subtotal
        })
        total += subtotal
        
    return render(request, 'order.html', {'items': items, 'total': total})

def order_create(request):
    if request.method == 'POST':
        name = request.POST.get('name')
        phone = request.POST.get('phone')
        address = request.POST.get('address')
        cart = request.session.get('cart', {})
        
        if cart:
            total = sum(item['price'] * item['quantity'] for item in cart.values())
            # Modelingizga mos maydonlarni ishlatamiz (first_name, last_name, total_price)
            order = Order.objects.create(
                first_name=name,
                last_name="", # Placeholder
                phone=phone,
                address=address,
                total_price=total,
                status='Pending'
            )
            # Create order items
            for p_id, item in cart.items():
                product = Product.objects.get(id=p_id)
                OrderItem.objects.create(
                    order=order,
                    product=product,
                    quantity=item['quantity'],
                    price=item['price']
                )
            request.session['cart'] = {}
            return JsonResponse({'success': True, 'order_id': order.id})
    return JsonResponse({'success': False})

from django.contrib.auth.decorators import login_required

# Admin Statistics & Dashboard
@login_required(login_url='login_page')
def home_page(request):
    now = timezone.now()
    today_start = now.replace(hour=0, minute=0, second=0, microsecond=0)
    week_start = today_start - timedelta(days=7)
    
    counts = {
        'orders': Order.objects.count(),
        'products': Product.objects.count(),
        'categories': Category.objects.count(),
        'revenue': Order.objects.aggregate(total=Sum('total_price'))['total'] or 0,
        'today_revenue': Order.objects.filter(created_at__gte=today_start).aggregate(total=Sum('total_price'))['total'] or 0,
        'weekly_revenue': Order.objects.filter(created_at__gte=week_start).aggregate(total=Sum('total_price'))['total'] or 0
    }
    
    # Charts data: Categories
    categories = Category.objects.annotate(p_count=Count('products'))
    cat_names = [c.name for c in categories]
    cat_counts = [c.p_count for c in categories]
    
    # Top Selling Products (Real-world metric)
    top_products = Product.objects.annotate(
        sell_count=Sum('orderitem__quantity')
    ).filter(sell_count__gt=0).order_by('-sell_count')[:5]
    
    # New orders alert (Last 1 hour)
    one_hour_ago = now - timedelta(hours=1)
    new_orders = Order.objects.filter(created_at__gte=one_hour_ago, status='Pending').count()
    
    # Recent Activities (Audit Log)
    last_activities = AuditLog.objects.all().order_by('-created_at')[:5]
    
    # Recent Orders (Today's)
    recent_orders = Order.objects.filter(created_at__gte=today_start).order_by('-created_at')[:5]
    
    context = {
        'counts': counts,
        'cat_names': json.dumps(cat_names),
        'cat_counts': json.dumps(cat_counts),
        'new_orders_alert': new_orders,
        'top_products': top_products,
        'last_activities': last_activities,
        'recent_orders': recent_orders
    }
    return render(request, 'index.html', context)

@login_required(login_url='login_page')
def check_new_orders(request):
    # Oxirgi 1 daqiqa ichida tushgan yangi buyurtmalar soni
    one_min_ago = timezone.now() - timedelta(minutes=1)
    count = Order.objects.filter(created_at__gte=one_min_ago, status='Pending').count()
    return JsonResponse({'new_orders': count})

from functools import wraps
from django.utils.text import slugify

def manager_required(view_func):
    @wraps(view_func)
    def _wrapped_view(request, *args, **kwargs):
        if request.user.is_superuser or request.user.groups.filter(name='Manager').exists():
            return view_func(request, *args, **kwargs)
        return HttpResponse("Sizda bu sahifaga kirish huquqi yo'q!", status=403)
    return _wrapped_view

@login_required(login_url='login_page')
@manager_required
def category_list(request):
    models = Category.objects.all()
    return render(request, 'category/list.html', {'models': models})

@login_required(login_url='login_page')
@manager_required
def category_create(request):
    if request.method == 'POST':
        name = request.POST.get('name')
        slug = request.POST.get('slug')
        if not slug:
            slug = slugify(name)
        cat = Category.objects.create(name=name, slug=slug)
        log_action(request, 'create', 'Category', cat.id, f"Yangi kategoriya yaratildi: {cat.name}")
        return redirect('category_list')
    return render(request, 'category/form.html')

@login_required(login_url='login_page')
@manager_required
def category_edit(request, pk):
    cat = get_object_or_404(Category, pk=pk)
    if request.method == 'POST':
        cat.name = request.POST.get('name')
        slug = request.POST.get('slug')
        if not slug:
            slug = slugify(cat.name)
        cat.slug = slug
        cat.save()
        log_action(request, 'edit', 'Category', cat.id, f"Kategoriya tahrirlandi: {cat.name}")
        return redirect('category_list')
    return render(request, 'category/form.html', {'category': cat})

@login_required(login_url='login_page')
@manager_required
def category_delete(request, pk):
    cat = get_object_or_404(Category, pk=pk)
    log_action(request, 'delete', 'Category', cat.id, f"Kategoriya o'chirildi: {cat.name}")
    cat.delete()
    return redirect('category_list')

# Other admin view placeholders
@login_required(login_url='login_page')
@manager_required
def product_admin_list(request):
    query = request.GET.get('q', '')
    products = Product.objects.all().select_related('category')
    if query:
        products = products.filter(
            Q(name__icontains=query) | 
            Q(description__icontains=query)
        )
    return render(request, 'product/list_admin.html', {'products': products, 'search_query': query})

@login_required(login_url='login_page')
@manager_required
def product_create(request):
    form = ProductForm()
    if request.method == 'POST':
        form = ProductForm(request.POST, request.FILES)
        if form.is_valid():
            p = form.save()
            log_action(request, 'create', 'Product', p.id, f"Yangi mahsulot qo'shildi: {p.name}")
            return redirect('product_admin_list')
    return render(request, 'product/form.html', {'form': form})

@login_required(login_url='login_page')
@manager_required
def product_edit(request, pk):
    product = get_object_or_404(Product, pk=pk)
    form = ProductForm(instance=product)
    if request.method == 'POST':
        form = ProductForm(request.POST, request.FILES, instance=product)
        if form.is_valid():
            p = form.save()
            log_action(request, 'edit', 'Product', p.id, f"Mahsulot tahrirlandi: {p.name}")
            return redirect('product_admin_list')
    return render(request, 'product/form.html', {'form': form, 'model': product})

@login_required(login_url='login_page')
@manager_required
def product_delete(request, pk):
    product = get_object_or_404(Product, pk=pk)
    log_action(request, 'delete', 'Product', product.id, f"Mahsulot o'chirildi: {product.name}")
    product.delete()
    return redirect('product_admin_list')

@login_required(login_url='login_page')
def order_admin_list(request):
    query = request.GET.get('q', '')
    status_filter = request.GET.get('status', '')
    start_date = request.GET.get('start_date', '')
    end_date = request.GET.get('end_date', '')
    
    orders = Order.objects.all().order_by('-created_at')
    
    if query:
        orders = orders.filter(
            Q(first_name__icontains=query) | 
            Q(phone__icontains=query) |
            Q(address__icontains=query) |
            Q(id__icontains=query)
        )
    
    if status_filter:
        orders = orders.filter(status=status_filter)
        
    if start_date:
        orders = orders.filter(created_at__date__gte=start_date)
    
    if end_date:
        orders = orders.filter(created_at__date__lte=end_date)
        
    context = {
        'orders': orders,
        'search_query': query,
        'status_filter': status_filter,
        'start_date': start_date,
        'end_date': end_date,
        'available_statuses': ['new', 'processing', 'completed', 'cancelled']
    }
    return render(request, 'order/list_admin.html', context)

def order_detail(request, pk):
    order = get_object_or_404(Order, pk=pk)
    return render(request, 'order/detail.html', {'order': order})

def order_status_update(request, pk, status):
    order = get_object_or_404(Order, pk=pk)
    order.status = status
    order.save()
    log_action(request, 'edit', 'Order', order.id, f"Buyurtma statusi o'zgartirildi: {status}")
    return redirect('order_admin_list')

def order_delete(request, pk):
    order = get_object_or_404(Order, pk=pk)
    log_action(request, 'delete', 'Order', order.id, f"Buyurtma o'chirildi: ID {order.id}")
    order.delete()
    return redirect('order_admin_list')

@login_required(login_url='login_page')
def export_orders_excel(request):
    import openpyxl
    from openpyxl.styles import Font, Alignment
    
    query = request.GET.get('q', '')
    status_filter = request.GET.get('status', '')
    start_date = request.GET.get('start_date', '')
    end_date = request.GET.get('end_date', '')
    
    orders = Order.objects.all().order_by('-created_at')
    
    if query:
        orders = orders.filter(
            Q(first_name__icontains=query) | Q(phone__icontains=query) | Q(id__icontains=query)
        )
    if status_filter:
        orders = orders.filter(status=status_filter)
    if start_date:
        orders = orders.filter(created_at__date__gte=start_date)
    if end_date:
        orders = orders.filter(created_at__date__lte=end_date)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Buyurtmalar"
    
    headers = ['ID', 'Mijoz', 'Tel', 'Manzil', 'Summa', 'Status', 'Sana']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
        
    for o in orders:
        ws.append([o.id, o.first_name, o.phone, o.address, float(o.total_price), o.status, o.created_at.strftime('%Y-%m-%d %H:%M')])
        
    log_action(request, 'export', 'Order', 0, f"Filtrlangan ({orders.count()} ta) buyurtmalar EXCELga yuklab olindi")
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=orders_filtered_{timezone.now().strftime("%Y%m%d_%H%M")}.xlsx'
    wb.save(response)
    return response

@login_required(login_url='login_page')
@manager_required
def branch_list(request):
    branches = Branch.objects.all()
    return render(request, 'branch/list.html', {'models': branches})

@login_required(login_url='login_page')
@manager_required
def branch_create(request):
    form = BranchForm()
    if request.method == 'POST':
        form = BranchForm(request.POST)
        if form.is_valid():
            b = form.save()
            log_action(request, 'create', 'Branch', b.id, f"Yangi filial qo'shildi: {b.name}")
            return redirect('branch_list')
    return render(request, 'branch/form.html', {'form': form})

@login_required(login_url='login_page')
@manager_required
def branch_edit(request, pk):
    branch = get_object_or_404(Branch, pk=pk)
    form = BranchForm(instance=branch)
    if request.method == 'POST':
        form = BranchForm(request.POST, instance=branch)
        if form.is_valid():
            form.save()
            log_action(request, 'edit', 'Branch', branch.id, f"Filial tahrirlandi: {branch.name}")
            return redirect('branch_list')
    return render(request, 'branch/form.html', {'form': form, 'model': branch})

@login_required(login_url='login_page')
@manager_required
def branch_delete(request, pk):
    branch = get_object_or_404(Branch, pk=pk)
    log_action(request, 'delete', 'Branch', branch.id, f"Filial o'chirildi: {branch.name}")
    branch.delete()
    return redirect('branch_list')

@login_required(login_url='login_page')
def audit_log_list(request):
    if not request.user.is_superuser:
        return HttpResponse("Faqat Super-Admin faolliklarni ko'ra oladi!", status=403)
    logs = AuditLog.objects.all().order_by('-created_at')[:100]
    return render(request, 'audit_logs.html', {'logs': logs})